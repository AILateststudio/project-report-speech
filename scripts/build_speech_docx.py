#!/usr/bin/env python3
"""Build a speech-script .docx from JSON. Usage: build_speech_docx.py in.json out.docx [--normal]"""
import json, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn

def build(data, out, compact=True):
    d = Document()
    size = 10.5 if compact else 12
    st = d.styles['Normal']
    st.font.size = Pt(size); st.font.name = '宋体'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    st.paragraph_format.line_spacing = 1.15 if compact else 1.5
    if compact:
        for s in d.sections:
            s.top_margin = s.bottom_margin = Cm(1.8)
            s.left_margin = s.right_margin = Cm(2)
    def para(text, bold=False, sz=size, color=None, after=1 if compact else 6):
        p = d.add_paragraph(); r = p.add_run(text); r.bold = bold; r.font.size = Pt(sz)
        if color: r.font.color.rgb = RGBColor(*color)
        p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
        return p
    t = para(data['title'], True, size + 2.5, after=4); t.alignment = 1
    for sec in data['sections']:
        h = para(f"【{sec['heading']}】", True, size, (0x1F, 0x4E, 0x79), 0)
        h.paragraph_format.keep_with_next = True
        for txt in sec['paragraphs']:
            para(txt)
    d.save(out)

if __name__ == '__main__':
    if len(sys.argv) < 3: sys.exit(__doc__)
    build(json.load(open(sys.argv[1], encoding='utf-8')), sys.argv[2], compact='--normal' not in sys.argv)
